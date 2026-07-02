"""
02324 离散数学 · 思维导图 生成器
从 .docx 笔记直接生成交互式 HTML 思维导图，零 LLM 依赖。

用法:
  python generate.py [--docx 02324_笔记.docx]

依赖:
  pip install python-docx
"""

import argparse
import os
import re
import sys
from collections import OrderedDict

from docx import Document
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.abspath(__file__))

# ── HTML 模板（复用已有样式） ──

CSS = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>德克离散数学 · 思维导图</title>
<style>
:root {
  --c1:#3b82f6; --c2:#10b981; --c3:#f59e0b;
  --c4:#ec4899; --c5:#8b5cf6; --c6:#06b6d4;
  --c7:#f97316; --c8:#84cc16; --c9:#e11d48;
  --bg:#0f172a; --card:#1e293b; --border:#334155;
  --text:#e2e8f0; --muted:#94a3b8;
}
* { box-sizing:border-box; margin:0; padding:0; }
body { background:var(--bg); color:var(--text); font-family:'PingFang SC','Microsoft YaHei',sans-serif; font-size:14px; }
header { text-align:center; padding:20px 16px 12px; }
header h1 { font-size:18px; font-weight:700; letter-spacing:1px; }
header p { color:var(--muted); font-size:12px; margin-top:4px; }
.chapters { padding:0 12px 40px; display:flex; flex-direction:column; gap:10px; }
.ch { border-radius:12px; overflow:hidden; border:1px solid var(--border); }
.ch-title { display:flex; align-items:center; justify-content:space-between; padding:13px 16px; cursor:pointer; user-select:none; font-weight:700; font-size:14px; transition:opacity .15s; }
.ch-title:active { opacity:.75; }
.ch-title .badge { font-size:11px; font-weight:400; opacity:.75; margin-left:6px; }
.ch-title .arrow { font-size:11px; transition:transform .25s; flex-shrink:0; }
.ch.open .arrow { transform:rotate(180deg); }
.ch1 .ch-title { background:linear-gradient(135deg,#1d3a8a,#1e40af); }
.ch2 .ch-title { background:linear-gradient(135deg,#064e3b,#065f46); }
.ch3 .ch-title { background:linear-gradient(135deg,#78350f,#92400e); }
.ch4 .ch-title { background:linear-gradient(135deg,#831843,#9d174d); }
.ch5 .ch-title { background:linear-gradient(135deg,#4c1d95,#5b21b6); }
.ch6 .ch-title { background:linear-gradient(135deg,#0e7490,#0369a1); }
.ch7 .ch-title { background:linear-gradient(135deg,#9a3412,#c2410c); }
.ch8 .ch-title { background:linear-gradient(135deg,#3f6212,#4d7c0f); }
.ch9 .ch-title { background:linear-gradient(135deg,#881337,#9f1239); }
.ch-body { background:var(--card); display:none; }
.ch.open .ch-body { display:block; }
.kp { border-bottom:1px solid var(--border); }
.kp:last-child { border-bottom:none; }
.kp-title { display:flex; align-items:center; gap:10px; padding:10px 16px; cursor:pointer; user-select:none; transition:background .15s; }
.kp-title:active { background:rgba(255,255,255,.04); }
.kp-num { font-size:11px; color:var(--muted); min-width:28px; }
.kp-name { flex:1; font-size:13px; }
.kp-arrow { font-size:10px; color:var(--muted); transition:transform .2s; }
.kp.open .kp-arrow { transform:rotate(180deg); }
.kp-detail { display:none; padding:10px 16px 14px 54px; font-size:12px; color:var(--muted); line-height:1.8; }
.kp.open .kp-detail { display:block; }
.kp-detail b { color:var(--text); }
.kp-detail .tag { display:inline-block; background:rgba(255,255,255,.08); border-radius:4px; padding:1px 6px; font-size:11px; margin:1px 2px; color:#cbd5e1; }
.kp-detail .formula { display:block; background:#0f172a; border-radius:6px; padding:6px 10px; margin:6px 0; font-size:12px; color:#93c5fd; letter-spacing:.5px; overflow-x:auto; white-space:pre-wrap; font-family:'Consolas','Courier New',monospace; }
.kp-detail table { width:100%; border-collapse:collapse; margin:6px 0; font-size:11px; }
.kp-detail td, .kp-detail th { border:1px solid var(--border); padding:4px 6px; text-align:center; }
.kp-detail th { background:rgba(255,255,255,.06); color:var(--text); }
.kp-detail ul { padding-left:14px; }
.kp-detail li { margin:2px 0; }
.kp-detail .tip { color:#fbbf24; }
.kp-detail .ok  { color:#34d399; }
.kp-detail .err { color:#f87171; }
.kp-detail .kp-sub { display:block; color:var(--text); font-size:12px; margin-top:8px; font-weight:600; }
.kp-detail .em { color:#fbbf24; font-weight:600; }
</style>
</head>
<body>
<header>
  <h1>德克离散数学 · 课程笔记</h1>
  <p>共 {ch_count} 章 · 点击章节/考点展开详情</p>
</header>
<div class="chapters">
"""

FOOTER = """
<script>
function toggleCh(el){var ch=el.closest('.ch'),isOpen=ch.classList.contains('open');ch.classList.toggle('open');if(!isOpen){ch.querySelectorAll('.kp-detail').forEach(function(d){d.style.display='none'});ch.querySelectorAll('.kp').forEach(function(k){k.classList.remove('open')})}}
function toggleKp(el){var kp=el.closest('.kp'),isOpen=kp.classList.contains('open');kp.classList.toggle('open');kp.querySelector('.kp-detail').style.display=isOpen?'none':'block'}
</script>
</div></body></html>
"""

# ── docx 解析 ──


def iter_block_items(doc):
    """按文档顺序遍历段落和表格。"""
    from docx.oxml.ns import qn
    body = doc.element.body
    para_idx = 0
    table_idx = 0
    for child in body:
        if child.tag == qn('w:p'):
            yield ('paragraph', doc.paragraphs[para_idx])
            para_idx += 1
        elif child.tag == qn('w:tbl'):
            if table_idx < len(doc.tables):
                yield ('table', doc.tables[table_idx])
                table_idx += 1


def table_to_html(table):
    """docx 表格 → HTML table 字符串。"""
    rows = []
    for ri, row in enumerate(table.rows):
        cells = []
        for ci, cell in enumerate(row.cells):
            tag = 'th' if ri == 0 else 'td'
            text = cell.text.strip()
            if not text and ri == 0:
                tag = 'th'
            cells.append(f'<{tag}>{h(text)}</{tag}>')
        rows.append(f'<tr>{"".join(cells)}</tr>')
    return f'<table>{"".join(rows)}</table>'


def h(text):
    """HTML 转义。"""
    return (text
            .replace('&', '&amp;')
            .replace('<', '&lt;')
            .replace('>', '&gt;')
            .replace('"', '&quot;'))


def format_content(text):
    """将纯文本转换为 HTML 格式。保留基本标记。"""
    text = h(text)
    # 保留常见符号
    text = text.replace('¬', '¬')
    text = text.replace('∧', '∧')
    text = text.replace('∨', '∨')
    text = text.replace('→', '→')
    text = text.replace('↔', '↔')
    text = text.replace('⇔', '⇔')
    text = text.replace('¬', '¬')
    text = text.replace('∈', '∈')
    text = text.replace('⊆', '⊆')
    text = text.replace('∪', '∪')
    text = text.replace('∩', '∩')
    text = text.replace('∅', '∅')
    text = text.replace('∀', '∀')
    text = text.replace('∃', '∃')
    # 加粗标记 **text** → <b>text</b>
    text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
    # 换行
    text = text.replace('\n', '<br>')
    return text


def docx_to_html(docx_path):
    """主解析函数：docx → HTML 字符串。"""
    doc = Document(docx_path)

    chapters = []  # [(chapter_title, [(kp_num, kp_title, [(sub_title, [content_lines])])])]
    current_chapter = None
    current_kp = None
    current_sub = None
    chapter_kp_count = 0
    kp_num = 0

    # 跳过标题和目录的状态
    skip_title = True  # 跳过第一个 Heading 1 (德克离散数学_课程笔记)
    seen_first_h1 = False

    # 临时缓冲：当一个 Heading 3 出现前的 Normal 段落归入上一个 Heading 3，
    # 在 Heading 2 前的 Normal 段落归入上个 Heading 2
    # 我们用累积模式
    pending_text = []  # 不属于任何 H3 的累积文本（归入当前 Kp）

    # pending_text 归入当前考点（作为无标题子主题）
    # 在切换考点或保存章节前刷新
    def flush_pending():
        nonlocal pending_text
        if pending_text and current_kp is not None:
            current_kp[1].append(('', pending_text.copy()))
            pending_text = []

    for kind, item in iter_block_items(doc):
        if kind == 'paragraph':
            p = item
            style_name = p.style.name
            text = p.text.strip()

            if not text:
                continue

            if style_name.startswith('Heading'):
                level = int(style_name.split()[-1])

                if level == 1:
                    if skip_title:
                        skip_title = False
                        continue
                    # 新章节
                    if current_chapter is not None:
                        flush_pending()
                        if current_kp is not None:
                            current_chapter.append((kp_num, current_kp[0], current_kp[1]))
                            current_kp = None
                        chapters.append((current_chapter_title, chapter_kp_count, current_chapter))
                    current_chapter = []
                    current_chapter_title = text
                    chapter_kp_count = 0
                    current_kp = None
                    current_sub = None
                    kp_num = 0
                    pending_text = []

                elif level == 2:
                    flush_pending()
                    # 考点
                    if current_kp is not None:
                        current_chapter.append((kp_num, current_kp[0], current_kp[1]))
                    kp_num += 1
                    chapter_kp_count += 1
                    current_kp = [text, []]  # [title, [(sub_title, content_lines)]]
                    current_sub = None
                    pending_text = []

                elif level == 3:
                    flush_pending()
                    # 子主题
                    if current_kp is None:
                        # 没有 H2 直接出现 H3，创建一个无名考点
                        kp_num += 1
                        chapter_kp_count += 1
                        current_kp = ['', []]
                    current_sub = text
                    current_kp[1].append((text, []))
                    pending_text = []

            else:
                # Normal 段落 → 内容
                if current_sub is not None and current_kp is not None and current_kp[1]:
                    # 属于当前 H3
                    current_kp[1][-1][1].append(text)
                elif current_kp is not None:
                    # 属于当前考点（没有 H3 的内容）
                    pending_text.append(text)
                # 如果既无考点也无子主题，忽略（如目录文本）

        elif kind == 'table':
            tbl_html = table_to_html(item)
            if current_sub is not None and current_kp is not None and current_kp[1]:
                current_kp[1][-1][1].append(('__table__', tbl_html))
            elif current_kp is not None:
                pending_text.append(('__table__', tbl_html))

    # 保存最后一章
    flush_pending()
    if current_kp is not None:
        current_chapter.append((kp_num, current_kp[0], current_kp[1]))
    if current_chapter:
        chapters.append((current_chapter_title, chapter_kp_count, current_chapter))

    # ── 渲染 HTML ──
    ch_index = 0
    ch_html_parts = []

    for ch_title, kp_count, kps in chapters:
        ch_index += 1
        ch_class = f'ch{ch_index}'
        # 从标题提取序号（第一章 → 1）
        seq_match = re.search(r'第[一二三四五六七八九十]+章', ch_title)
        if seq_match:
            cn = seq_match.group()
            num_map = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5,
                       '六': 6, '七': 7, '八': 8, '九': 9, '十': 10}
            for k, v in num_map.items():
                if k in cn:
                    ch_class = f'ch{v}'
                    break

        parts = [f'<div class="ch {ch_class}">']
        title_display = ch_title.replace('：', ' · ', 1) if '：' in ch_title else ch_title
        # 去掉可能的数字序号前缀
        title_display = re.sub(r'^\d+[\.、]\s*', '', title_display)
        parts.append(
            f'<div class="ch-title" onclick="toggleCh(this)">'
            f'<span>{h(title_display)}<span class="badge">{kp_count}考点</span></span>'
            f'<span class="arrow">▼</span></div>'
        )
        parts.append('<div class="ch-body">')

        for kp_num_item, kp_title, subs in kps:
            # 解析考点序号
            num_match = re.match(r'考点(\d+)', kp_title)
            if num_match:
                display_num = f'考点{num_match.group(1)}'
                name = kp_title.split('：', 1)[-1] if '：' in kp_title else kp_title
            else:
                display_num = f'考点{kp_num_item}'
                name = kp_title if kp_title else '(未命名)'

            parts.append(
                f'<div class="kp">'
                f'<div class="kp-title" onclick="toggleKp(this)">'
                f'<span class="kp-num">{display_num}</span>'
                f'<span class="kp-name">{h(name)}</span>'
                f'<span class="kp-arrow">▼</span></div>'
                f'<div class="kp-detail">'
            )

            # 子主题
            for sub_title, content_lines in subs:
                if sub_title:
                    parts.append(f'<span class="kp-sub">{h(sub_title)}</span>')
                for line in content_lines:
                    if isinstance(line, tuple) and line[0] == '__table__':
                        parts.append(line[1])
                    elif isinstance(line, str):
                        line = line.strip()
                        if not line:
                            continue
                        # 检测公式（含 →, ∨, ∧, ¬ 等符号的行）
                        if any(sym in line for sym in ['→', '↔', '⇔', '∨', '∧', '¬', '∀', '∃']):
                            parts.append(f'<span class="formula">{format_content(line)}</span>')
                        else:
                            out = format_content(line)
                            if '<br>' in out:
                                parts.append(out)
                            else:
                                parts.append(f'{out}<br>')

            # 如果有 pending_text（属于考点但不属于任何子主题），放在子主题之后
            # 但我们已经把 pending_text 作为内容加入了

            parts.append('</div></div>')

        parts.append('</div></div>')
        ch_html_parts.append('\n'.join(parts))

    chapter_html = '\n\n'.join(ch_html_parts)

    # 组装完整 HTML
    html = CSS.replace('{ch_count}', str(ch_index)) + '\n' + chapter_html + '\n' + FOOTER
    return html


def main():
    parser = argparse.ArgumentParser(description='从 docx 生成思维导图 HTML')
    parser.add_argument('--docx', default=None, help='.docx 文件路径（默认自动查找）')
    parser.add_argument('-o', '--output', default=None, help='输出 HTML 路径（默认 index.html）')
    args = parser.parse_args()

    # 查找 docx
    docx_path = args.docx
    if docx_path is None:
        candidates = [f for f in os.listdir(ROOT) if f.endswith('.docx') and not f.startswith('~')]
        if not candidates:
            print('❌ 未找到 .docx 文件')
            sys.exit(1)
        docx_path = os.path.join(ROOT, candidates[0])
        print(f'📄 使用: {os.path.basename(docx_path)}')

    output = args.output or os.path.join(ROOT, 'index.html')

    print(f'📖 正在解析: {docx_path}')
    html = docx_to_html(docx_path)

    with open(output, 'w', encoding='utf-8') as f:
        f.write(html)

    file_size = os.path.getsize(output)
    line_count = html.count('\n')
    print(f'✅ 生成成功: {output}')
    print(f'   大小: {file_size/1024:.1f} KB, {line_count} 行')

    # 验证 HTML 结构
    div_open = html.count('<div')
    div_close = html.count('</div>')
    if div_open == div_close:
        print(f'   ✅ div 标签平衡: {div_open}')
    else:
        print(f'   ⚠️  div 标签不平衡: open={div_open} close={div_close}')


if __name__ == '__main__':
    main()
